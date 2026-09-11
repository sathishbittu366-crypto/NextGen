import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_import import match_headers, roll_prefix_to_batch
from sms_app.services import learning_service as ls


class FakeCursor:
    def __init__(self, students):
        self.lastrowid = 9001
        self.students = students
        self.items = []

    def execute(self, sql, params=()):
        normalized = " ".join(sql.lower().split())
        self._last_sql = normalized
        self._last_params = params
        if normalized.startswith("select id,code,name from academic_semesters"):
            return self
        if normalized.startswith("select 1 from students where department=? and batch=? limit 1"):
            return self
        if normalized.startswith("select roll_no,name from students"):
            return self
        if normalized.startswith("insert into result_batches"):
            return self
        if normalized.startswith("insert into result_items"):
            self.items.append(params)
            return self
        return self

    def fetchone(self):
        sql = getattr(self, "_last_sql", "")
        params = getattr(self, "_last_params", ())
        if sql.startswith("select 1 from students where department=? and batch=? limit 1"):
            return {"1": 1} if len(params) == 2 and params[1] == "2024-2028" else None
        if sql.startswith("select roll_no,name from students"):
            return {"roll_no": params[0], "name": "Student"} if len(params) == 3 and params[0] in self.students and params[2] == "2024-2028" else None
        return {"id": 3, "code": "II-I", "name": "II B.Tech I Semester"}


class FakeConnection:
    def __init__(self, students):
        self.cursor = FakeCursor(students)

    def __enter__(self):
        return self.cursor

    def __exit__(self, *_):
        return False


def workbook_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def test_short_serial_header_cannot_steal_exact_roll_header():
    report = match_headers(
        ["S NO", "H T NO"],
        ls.RESULT_ROLL_SPECS,
    )
    assert report.as_dict()["ignored"] == ["S NO"]
    assert report.as_dict()["mapped"] == [
        {"header": "H T NO", "field": "roll_no", "matched_via": "exact"}
    ]


def test_wide_results_upload_and_mapping(monkeypatch):
    students = {"24BT1A6701", "24BT1A6702"}
    fake = FakeConnection(students)
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)

    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None, None, "24CS302PC", None, None, None, None, None],
        [None, None, "Digital Electronics", None, None, None, None, None, "Data Structures", None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C", "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6701", 36, 48, 84, "A+", 9, 3, 36, 37, 73, "A", 8, 3],
        [2, "24BT1A6702", 38, 44, 82, "A", 8, 3, 39, 35, 74, "A", 8, 3],
    ])

    result = ls.upload_results_excel(
        raw=raw,
        filename="wide.xlsx",
        department="CSD",
        batch="2024-2028",
        semester_id=3,
        title="Result",
        admin_username="admin",
    )

    assert result["source_format"] == "wide"
    assert result["rows_imported"] == 4
    assert result["students_affected"] == 2
    assert len(fake.cursor.items) == 4
    first = fake.cursor.items[0]
    assert first[6:9] == (36.0, 48.0, 3.0)
    assert any(x["field"].startswith("subject_block:24CS301PC") for x in result["column_mapping"]["mapped"])
    assert "S NO" in result["column_mapping"]["ignored"]


def test_long_results_path_survives_and_reports_unknown_columns(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)

    raw = workbook_bytes([
        ["Rollnumber", "SubjectCode", "Subject Name", "Marks", "Unknown Thing"],
        ["24BT1A6701", "24CS301PC", "Digital Electronics", 84, "ignored value"],
    ])
    result = ls.upload_results_excel(
        raw=raw,
        filename="long.xlsx",
        department="CSD",
        batch="2024-2028",
        semester_id=3,
        title="Long",
        admin_username="admin",
    )

    assert result["source_format"] == "long"
    assert result["rows_imported"] == 1
    assert "Unknown Thing" in result["column_mapping"]["ignored"]
    assert next(x for x in result["column_mapping"]["mapped"] if x["header"] == "Rollnumber")["matched_via"] == "fuzzy"


def test_wide_import_fails_closed_on_invalid_numeric_value(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None, None],
        [None, None, "Digital Electronics", None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6701", 36, "NOT-A-NUMBER", 84, "A+", 9, 3],
    ])
    with pytest.raises(ValueError, match=r"Row 4: external marks \(EM\) must be numeric"):
        ls.upload_results_excel(
            raw=raw, filename="bad-wide.xlsx", department="CSD", batch="2024-2028", semester_id=3,
            title="Result", admin_username="admin",
        )
    assert fake.cursor.items == []


def test_wide_results_handles_trailing_summary_columns(monkeypatch):
    # — WHY: §7.1 regression guard. Real VR24 sheets append single-column
    # SGPA/Credits/Backlogs fields after the last 6-column subject block —
    # these are NOT a malformed extra subject block and must not crash the
    # col += 6 block-walking loop. See RESULT_SUMMARY_FIELD_SPECS.
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None, None, "SGPA", "Credits", "Backlogs"],
        [None, None, "Digital Electronics", None, None, None, None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C", None, None, None],
        [1, "24BT1A6701", 36, 48, 84, "A+", 9, 3, "8.82", "3", "0"],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="trailing.xlsx", department="CSD", batch="2024-2028", semester_id=3,
        title="Result", admin_username="admin",
    )
    assert result["rows_imported"] == 1
    item = fake.cursor.items[0]
    assert item[12] == "8.82"  # sgpa broadcast onto the subject row
    mapped_fields = {m["field"] for m in result["column_mapping"]["mapped"]}
    assert "summary:sgpa" in mapped_fields
    assert "summary:credits_total" in mapped_fields
    assert "summary:backlogs" in mapped_fields


def test_wide_results_blank_im_em_credits_store_as_null(monkeypatch):
    # — WHY: §7.2 regression guard. Real VR24 sheets legitimately leave
    # IM/EM/Credits blank for some subjects (e.g. an internal-only lab).
    # A blank means "no such component" and must store as SQL NULL — never
    # coerced to 0, which would falsely claim a scored value.
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24MC309", None, None, None, None, None],
        [None, None, "Gender Sensitization Lab", None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6701", 45, None, 45, "B", 6, None],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="blank-im-em.xlsx", department="CSD", batch="2024-2028", semester_id=3,
        title="Result", admin_username="admin",
    )
    assert result["rows_imported"] == 1
    item = fake.cursor.items[0]
    assert item[6] == 45.0   # internal_marks
    assert item[7] is None   # external_marks: blank -> NULL, not 0
    assert item[8] is None   # credits: blank -> NULL, not 0


def test_wide_results_stops_at_footer_rows_after_real_students(monkeypatch):
    # — WHY: §7.3 regression guard. VR24 sheets bury a REGISTERED/APPEARED/
    # PASSED/FAILED/PASS%/FACULTY/SIGNATURE statistics-and-sign-off footer
    # directly after the last real student row. The row loop must stop at
    # the first non-blank row whose roll-number cell isn't roll-shaped,
    # rather than trying to look up "REGISTERED" etc. as a student.
    fake = FakeConnection({"24BT1A6701", "25BT5A6703"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None, None],
        [None, None, "Digital Electronics", None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6701", 36, 48, 84, "A+", 9, 3],
        [2, "25BT5A6703", 27, 17, 44, "F", 0, 0],   # lateral-entry roll shape
        [None, None, None, None, None, None, None, None],  # blank separator
        [None, "REGISTERED", 2, None, None, None, None, None],
        [None, "APPEARED", 2, None, None, None, None, None],
        [None, "PASS %", 100, None, None, None, None, None],
        [None, "FACULTY", "Dr. Rao", None, None, None, None, None],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="with-footer.xlsx", department="CSD", batch="2024-2028", semester_id=3,
        title="Result", admin_username="admin",
    )
    assert result["rows_imported"] == 2
    assert result["students_affected"] == 2
    rolls_imported = {item[1] for item in fake.cursor.items}
    assert rolls_imported == {"24BT1A6701", "25BT5A6703"}


def test_wide_results_ab_marker_stores_null_external_marks_not_zero(monkeypatch):
    # — WHY: real VR24 mark sheets use "AB" (Absent) as a distinct, meaningful
    # marker in EM/G — not a blank, not garbage. Confirmed against the real
    # reference file: every occurrence is internally consistent (EM='AB',
    # TM=IM exactly, G='AB', GP=0, C=0). Per the project owner: AB means "did
    # not attempt" and must never be stored as 0 ("attempted and scored
    # zero") — external_marks must be NULL, and grade must read "AB".
    fake = FakeConnection({"24BT1A6721"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None, None],
        [None, None, "Digital Electronics", None, None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6721", 22, "AB", 22, "AB", 0, 0],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="absent.xlsx", department="CSD", batch="2024-2028", semester_id=3,
        title="Result", admin_username="admin",
    )
    assert result["rows_imported"] == 1
    item = fake.cursor.items[0]
    assert item[6] == 22.0     # internal_marks: real value, unaffected
    assert item[7] is None    # external_marks: AB -> NULL, never 0
    assert item[9] == "AB"    # grade: forced to the literal marker


def test_templates_match_required_shapes_and_are_real_xlsx():
    students = load_workbook(io.BytesIO(ls.build_students_template()), data_only=True)
    results = load_workbook(io.BytesIO(ls.build_results_template()), data_only=True)
    ws = results.active
    assert [ws.cell(3, c).value for c in range(3, 9)] == ["IM", "EM", "TM", "G", "GP", "C"]
    assert ws.cell(1, 3).value == "24CS301PC"
    assert ws.cell(2, 3).value == "Digital Electronics"
    assert ws.cell(4, 2).value == "24CSDS0001"
    assert students.active.max_row >= 2
    students.close(); results.close()


def test_database_contract_contains_additive_result_columns():
    text = Path("database.py").read_text()
    for name, sql_type in [
        ("internal_marks", "DECIMAL(10,2) NULL"),
        ("external_marks", "DECIMAL(10,2) NULL"),
        ("credits", "DECIMAL(5,2) NULL"),
    ]:
        assert f"{name} {sql_type}" in text
        assert f'if column not in existing_result_item_cols' in text
    assert "SHOW COLUMNS FROM result_items" in text


def test_roll_prefix_to_batch_is_stable_and_fails_closed():
    assert roll_prefix_to_batch("24BT1A6701") == "2024-2028"
    assert roll_prefix_to_batch("26BC1A0001") == "2026-2030"
    assert roll_prefix_to_batch("X") is None
    assert roll_prefix_to_batch("171234") is None
    assert roll_prefix_to_batch("361234") is None

def test_student_batch_fallback_uses_live_current_year(monkeypatch):
    import api.routes_students as rs

    class FrozenDateTime:
        @classmethod
        def now(cls):
            return cls()

        @property
        def year(self):
            return 2031

    monkeypatch.setattr(rs, "datetime", FrozenDateTime)

    assert rs._compute_year_and_batch({"roll_no": "MALFORMED", "current_semester_id": 1}) == ("1st Year", "2031-2035 Batch")
    assert rs._compute_year_and_batch({"roll_no": "MALFORMED", "current_semester_id": 3}) == ("2nd Year", "2030-2034 Batch")
    assert rs._compute_year_and_batch({"roll_no": "MALFORMED", "current_semester_id": 5}) == ("3rd Year", "2029-2033 Batch")
    assert rs._compute_year_and_batch({"roll_no": "MALFORMED", "current_semester_id": 7}) == ("4th Year", "2028-2032 Batch")
    # A parseable roll prefix remains the stable cohort identity.
    assert rs._compute_year_and_batch({"roll_no": "24BT1A6701", "current_semester_id": 7}) == ("4th Year", "2024-2028 Batch")


class CohortFixtureCursor(FakeCursor):
    def __init__(self):
        super().__init__({"24BT1A6701"})

    def execute(self, sql, params=()):
        normalized = " ".join(sql.lower().split())
        if normalized.startswith("select roll_no,name from students"):
            # Fixture student is currently in semester 5, but its stable cohort
            # is 2024-2028 and the upload below targets semester 3.
            if "current_semester_id" in normalized:
                return None if params[2] == 3 else self
            if "batch=?" in normalized:
                return self if params[2] == "2024-2028" else None
        return super().execute(sql, params)


def test_progressed_student_old_semester_key_fails_but_cohort_key_succeeds(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    fake.cursor = CohortFixtureCursor()
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)

    # This is the exact pre-change eligibility query: it cannot find a student
    # who has progressed from semester 3 to semester 5.
    assert fake.cursor.execute(
        "SELECT roll_no,name FROM students WHERE roll_no=? AND department=? AND current_semester_id=? AND active=1",
        ("24BT1A6701", "CSD", 3),
    ) is None

    raw = workbook_bytes([
        ["Roll Number", "Subject Code", "Subject Name", "Marks"],
        ["24BT1A6701", "24CS301PC", "Digital Electronics", 84],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="progressed.xlsx", department="CSD",
        batch="2024-2028", semester_id=3, title="Result", admin_username="admin",
    )
    assert result["rows_imported"] == 1


def test_results_upload_rejects_unknown_cohort_batch(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["Roll Number", "Subject Code", "Subject Name", "Marks"],
        ["24BT1A6701", "24CS301PC", "Digital Electronics", 84],
    ])
    with pytest.raises(ValueError, match="Selected batch does not exist for this branch"):
        ls.upload_results_excel(
            raw=raw, filename="unknown-batch.xlsx", department="CSD",
            batch="2099-2103", semester_id=3, title="Result", admin_username="admin",
        )


class ResultsReadCursor:
    def __init__(self):
        self.calls = []

    def execute(self, sql, params=()):
        normalized = " ".join(sql.lower().split())
        self.calls.append((normalized, params))
        if normalized.startswith("select roll_no,name,department,batch from students"):
            return self
        if normalized.startswith("select rb.id,rb.title,rb.created_at,rb.source_filename,rb.semester_id"):
            return self
        if normalized.startswith("select subject_code,subject_name,marks,max_marks,internal_marks"):
            return self
        return self

    def fetchone(self):
        # First fetch is the student. Subsequent upload-batch fetches are not
        # used by this test cursor; use fetchall for the multi-row reads below.
        return {"roll_no": "24BT1A6701", "name": "Alice", "department": "CSD", "batch": "2024-2028"}

    def fetchall(self):
        sql = self.calls[-1][0]
        if sql.startswith("select rb.id,rb.title,rb.created_at,rb.source_filename,rb.semester_id"):
            return [
                {"id": 11, "title": "II-I Result", "created_at": "2026-03-10", "source_filename": "ii-i.xlsx", "semester_id": 3, "semester_code": "II-I", "semester_name": "II B.Tech I Semester"},
                {"id": 12, "title": "II-II Result", "created_at": "2026-07-10", "source_filename": "ii-ii.xlsx", "semester_id": 4, "semester_code": "II-II", "semester_name": "II B.Tech II Semester"},
            ]
        if sql.startswith("select subject_code,subject_name,marks,max_marks,internal_marks"):
            batch_id = self.calls[-1][1][0]
            return [{
                "subject_code": "24CS301PC", "subject_name": "Digital Electronics", "marks": 84,
                "max_marks": 100, "internal_marks": 36, "external_marks": 48, "credits": 3,
                "grade": "A+", "grade_point": "9", "result_status": "PASS", "sgpa": "8.82", "percentage": "84",
            }] if batch_id == 11 else [{
                "subject_code": "24CS401PC", "subject_name": "Data Structures", "marks": 88,
                "max_marks": 100, "internal_marks": 38, "external_marks": 50, "credits": 3,
                "grade": "A+", "grade_point": "9", "result_status": "PASS", "sgpa": "9.01", "percentage": "88",
            }]
        return []


def test_student_results_reads_all_semesters_by_cohort(monkeypatch):
    fake = ResultsReadCursor()
    class Conn:
        def __enter__(self): return fake
        def __exit__(self, *_): return False
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: Conn())

    result = ls.get_student_results(roll_no="24BT1A6701")

    assert result["student"]["roll_no"] == "24BT1A6701"
    assert [entry["batch"]["semester_code"] for entry in result["results"]] == ["II-I", "II-II"]
    assert result["results"][0]["sgpa"] == "8.82"
    assert result["results"][1]["sgpa"] == "9.01"
    batch_lookup = next(params for sql, params in fake.calls if sql.startswith("select rb.id,rb.title,rb.created_at,rb.source_filename,rb.semester_id"))
    assert batch_lookup == ("CSD", "2024-2028")
    assert all("current_semester_id" not in sql for sql, _ in fake.calls)


def test_unregistered_students_are_skipped_in_wide_upload(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["S NO", "H T NO", "24CS301PC", None, None, None, None],
        [None, None, "Digital Electronics", None, None, None, None],
        [None, None, "IM", "EM", "TM", "G", "GP", "C"],
        [1, "24BT1A6701", 36, 48, 84, "A+", 9, 3],
        [2, "24BT1A6721", 30, 40, 70, "B", 7, 3],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="partial-wide.xlsx", department="CSD", batch="2024-2028",
        semester_id=3, title="Result", admin_username="admin",
    )

    assert result["rows_imported"] == 1
    assert result["students_affected"] == 1
    assert result["skipped_count"] == 1
    assert result["skipped_students"] == [{
        "row": 5,
        "roll_no": "24BT1A6721",
        "reason": "Student is not registered in the app for this branch/batch",
    }]
    assert [item[1] for item in fake.cursor.items] == ["24BT1A6701"]


def test_unregistered_students_are_skipped_in_long_upload(monkeypatch):
    fake = FakeConnection({"24BT1A6701"})
    monkeypatch.setattr(ls, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ls, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["Roll Number", "Subject Code", "Subject Name", "Marks"],
        ["24BT1A6701", "24CS301PC", "Digital Electronics", 84],
        ["24BT1A6721", "24CS301PC", "Digital Electronics", 70],
    ])
    result = ls.upload_results_excel(
        raw=raw, filename="partial-long.xlsx", department="CSD", batch="2024-2028",
        semester_id=3, title="Result", admin_username="admin",
    )

    assert result["rows_imported"] == 1
    assert result["skipped_count"] == 1
    assert result["skipped_students"][0]["roll_no"] == "24BT1A6721"
    assert [item[1] for item in fake.cursor.items] == ["24BT1A6701"]
